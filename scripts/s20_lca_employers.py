#!/usr/bin/env python3
"""S20-B instrument #3 — LCA employer-universe extractor.

Parses a DOL LCA disclosure xlsx into the unique-employer table:
  ingest/data/ats_seed/s20_census/lca_employers_fy2026q3.json

Usage:  python3 scripts/s20_lca_employers.py <path-to-xlsx> [out.json]

Canonical runtime = the s20-lca-employers.yml workflow on the org repo
(GHA US egress — the only transport that serves the full file; the HK
sandbox is Akamai-blocked and the supabase proxy truncates at ~10.5MB,
which is BELOW the real quarterly file size).
"""
from __future__ import annotations

import json
import statistics
import sys
import time
from collections import defaultdict
from pathlib import Path

HERE = Path(__file__).resolve().parent
REPO = HERE.parent
OUT_DEFAULT = REPO / "ingest/data/ats_seed/s20_census/lca_employers_fy2026q3.json"

FIELDS = {
    "CASE_NUMBER": "n", "CASE_STATUS": "st", "EMPLOYER_NAME": "em",
    "JOB_TITLE": "jt", "WORKSITE_CITY": "city", "WORKSITE_STATE": "wst",
    "WAGE_RATE_OF_PAY_FROM": "wage", "WAGE_UNIT_OF_PAY": "wu",
    "EMPLOYER_CITY": "ecity", "EMPLOYER_STATE": "est",
    "EMPLOYER_COUNTRY": "ecountry", "NAICS_CODE": "naics",
}


def main():
    xlsx = Path(sys.argv[1] if len(sys.argv) > 1 else
                "ingest/data/ats_seed/s20_census/LCA_Disclosure_Data_FY2026_Q3.xlsx")
    out = Path(sys.argv[2] if len(sys.argv) > 2 else OUT_DEFAULT)
    out.parent.mkdir(parents=True, exist_ok=True)
    from openpyxl import load_workbook
    wb = load_workbook(xlsx, read_only=True, data_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    header = [str(c or "").strip() for c in next(rows_iter)]
    idx = {}
    for k, alias in FIELDS.items():
        if k in header:
            idx[alias] = header.index(k)
    print("fields matched:", sorted(idx))
    employers = defaultdict(lambda: {
        "filings": 0, "states": defaultdict(int), "cities": set(),
        "titles": defaultdict(int), "wages": [], "status": defaultdict(int),
        "hq": "", "naics": "", "country": ""})
    n_rows = 0
    t0 = time.time()
    for row in rows_iter:
        n_rows += 1
        if "em" not in idx or not row[idx["em"]]:
            continue
        key = str(row[idx["em"]]).strip().upper()
        e = employers[key]
        e["filings"] += 1
        if "wst" in idx and row[idx["wst"]]:
            e["states"][str(row[idx["wst"]]).strip()] += 1
        if "city" in idx and row[idx["city"]]:
            if len(e["cities"]) < 30:
                e["cities"].add(str(row[idx["city"]]).strip().title())
        if "jt" in idx and row[idx["jt"]]:
            t = str(row[idx["jt"]]).strip()[:80]
            if len(e["titles"]) < 30:
                e["titles"][t] += 1
        if "st" in idx and row[idx["st"]]:
            e["status"][str(row[idx["st"]]).strip()] += 1
        if "wage" in idx and row[idx["wage"]] and "wu" in idx:
            try:
                v = float(row[idx["wage"]])
                unit = str(row[idx["wu"]] or "").lower()
                if unit in ("year", "yr", "annually", "annual") and \
                        10000 < v < 2000000 and len(e["wages"]) < 400:
                    e["wages"].append(v)
            except (TypeError, ValueError):
                pass
        if not e["hq"] and "ecity" in idx and row[idx["ecity"]]:
            st = str(row[idx["est"]]).strip() if "est" in idx and row[idx["est"]] else ""
            e["hq"] = f"{str(row[idx['ecity']]).title()}{', ' + st if st else ''}"
        if not e["country"] and "ecountry" in idx and row[idx["ecountry"]]:
            e["country"] = str(row[idx["ecountry"]]).strip()
        if not e["naics"] and "naics" in idx and row[idx["naics"]]:
            e["naics"] = str(row[idx["naics"]]).strip()
    print(f"rows: {n_rows}, unique employers: {len(employers)} "
          f"({time.time()-t0:.0f}s)")
    out_rows = []
    for name, e in employers.items():
        med = statistics.median(e["wages"]) if e["wages"] else None
        out_rows.append({
            "employer": name, "filings": e["filings"],
            "states": dict(sorted(e["states"].items(), key=lambda x: -x[1])),
            "top_titles": dict(sorted(e["titles"].items(),
                                      key=lambda x: -x[1])[:8]),
            "median_wage": round(med) if med else None,
            "status": dict(e["status"]), "hq": e["hq"],
            "naics": e["naics"], "country": e["country"],
        })
    out_rows.sort(key=lambda x: -x["filings"])
    (out).write_text(json.dumps(
        {"source": f"DOL LCA Disclosure {xlsx.stem} (GHA US egress)",
         "fetched": time.time(), "rows": n_rows,
         "unique_employers": len(out_rows), "employers": out_rows},
        indent=1))
    print("wrote:", out, out.stat().st_size, "bytes")
    for e in out_rows[:15]:
        print(f"  {e['filings']:>5}  {e['employer'][:55]}  "
              f"{list(e['states'].keys())[:4]}")


if __name__ == "__main__":
    main()
