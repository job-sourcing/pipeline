#!/usr/bin/env python3
"""DOL H-1B / LCA quarterly-extract builder (S11, design S8-F-3).

Downloads the DOL's quarterly LCA disclosure xlsx files
(https://www.dol.gov/agencies/eta/foreign-labor/performance — public
government data, no auth) and extracts the EMPLOYER-matching rows into
an append-only, case-number-deduped JSONL that board_dump's finish
phase joins per (normalized title, state) into the CSV wage-band
columns (v2.6 #44-#48).

WHY A STANDALONE SCRIPT (not a board_dump phase): the cadence differs
— the board refreshes weekly, the LCA files publish quarterly; and the
egress requirements differ — DOL's Akamai geo-blocks the HK sandbox
AND the Netlify US function (measured 2026-09-17: direct 403, Netlify
fetch-engine 403), while the Supabase Deno proxy PASSES but truncates
responses at ~10.5MB (the real files are larger — a "200" with a
truncated body is the trap this script's size-check catches). The one
egress measured to work for the full files is a plain US runner:
GITHUB ACTIONS (Azure US) — so this script's canonical runtime is the
`h1b-extract.yml` workflow on the org repo (free public minutes,
US egress, no size cap), committing the extract as its checkpoint.
The sandbox can still run it opportunistically (`--transport auto`
tries direct → curl_cffi impersonate → supabase-with-size-check).

Usage:
  python3 scripts/h1b_extract.py --recent 6            # last 6 published quarters
  python3 scripts/h1b_extract.py --quarters FY2025_Q4,FY2026_Q1
  python3 scripts/h1b_extract.py --probe               # sizes + egress check only

Output: {out-dir}/{label}.h1b_lca.jsonl — one line per filing:
  {caseNumber, caseStatus, caseSubmitted, decisionDate, employerName,
   jobTitle, socCode, socTitle, fullTimePosition, wageFrom, wageTo,
   wageUnit, prevailingWage, pwUnit, worksiteCity, worksiteState,
   worksitePostalCode, sourceFile}
"""
from __future__ import annotations

import argparse
import json
import os
import re
import sys
import time
from pathlib import Path
from urllib.parse import urlencode

HERE = Path(__file__).resolve().parent
REPO = HERE.parent
sys.path.insert(0, str(REPO / "ingest"))

DOL_PAGE = ("https://www.dol.gov/agencies/eta/foreign-labor/"
            "performance")
_FILE_RE = re.compile(r"LCA_Disclosure_Data_FY(\d{4})_Q(\d)\.xlsx")

# xlsx header → JSONL field. Tolerant of the DOL's occasional column
# renames: unknown headers are skipped, missing ones ship "".
_FIELDS = [
    ("CASE_NUMBER", "caseNumber"),
    ("CASE_STATUS", "caseStatus"),
    ("CASE_SUBMITTED", "caseSubmitted"),
    ("DECISION_DATE", "decisionDate"),
    ("EMPLOYER_NAME", "employerName"),
    ("JOB_TITLE", "jobTitle"),
    ("SOC_CODE", "socCode"),
    ("SOC_TITLE", "socTitle"),
    ("FULL_TIME_POSITION", "fullTimePosition"),
    ("WAGE_RATE_OF_PAY_FROM", "wageFrom"),
    ("WAGE_RATE_OF_PAY_TO", "wageTo"),
    ("WAGE_UNIT_OF_PAY", "wageUnit"),
    ("PREVAILING_WAGE", "prevailingWage"),
    ("PW_UNIT_OF_PAY", "pwUnit"),
    ("WORKSITE_CITY", "worksiteCity"),
    ("WORKSITE_STATE", "worksiteState"),
    ("WORKSITE_POSTAL_CODE", "worksitePostalCode"),
]


def _load_env():
    from jobsearch.config import load_config
    return load_config()


def _list_quarters(cfg) -> list[str]:
    """Published FY_Q quarters off the DOL performance page (via the
    supabase proxy — the page itself is small, well under the cap)."""
    import requests
    q = urlencode({"url": DOL_PAGE, "mode": "raw"})
    r = requests.get(
        f"{cfg.supabase_proxy_url}?{q}",
        headers={"Authorization": f"Bearer {cfg.supabase_proxy_token}",
                 "x-region": "us-east-1"},
        timeout=60)
    r.raise_for_status()
    found = sorted(set(f"FY{m.group(1)}_Q{m.group(2)}"
                       for m in _FILE_RE.finditer(r.text)),
                   key=lambda s: (s.split("_")[0], s.split("_Q")[1]))
    return found


def _file_url(quarter: str, alt: bool = False) -> str:
    """Canonical /sites/ path; `alt=True` = the /media/ path (some
    quarters publish ONLY there — FY2026_Q3 live-measured: /sites/ 404s
    while /media/ serves the file; the performance page links it with a
    double slash, which both forms tolerate)."""
    if alt:
        return (f"https://www.dol.gov/media/"
                f"LCA_Disclosure_Data_{quarter}.xlsx")
    return (f"https://www.dol.gov/sites/dolgov/files/ETA/oflc/pdfs/"
            f"LCA_Disclosure_Data_{quarter}.xlsx")


def _download(url: str, transport: str, cfg, timeout: int = 600
              ) -> tuple[bytes, str]:
    """Full-body download with transport fallback + TRUNCATION GUARD.

    The supabase proxy serves 200 with a silently truncated body at
    ~10.5MB — a body whose length disagrees with the target's
    Content-Length is an error, not a payload (that guard is what makes
    the fallback chain safe).

    Error contract: FileNotFoundError when ANY transport saw a
    definitive 404 (the path may simply be wrong — the caller's
    alt-path retry keys off this; run-4 lesson: wrapping the 404s into
    a blanket RuntimeError made that retry dead code. A 403 from ONE
    transport is Akamai noise, not a path verdict — the 404s from the
    transports that DID get an answer carry the verdict);
    RuntimeError when no transport got a definitive answer."""
    errors: list[str] = []
    saw_404 = False
    order = (["direct", "impersonate", "supabase"]
             if transport == "auto" else [transport])
    for t in order:
        try:
            body, clen = _fetch_via(t, url, cfg, timeout)
            if clen and len(body) != clen:
                errors.append(
                    f"{t}: TRUNCATED ({len(body)} of {clen} bytes)")
                continue
            return body, t
        except FileNotFoundError:
            errors.append(f"{t}: 404 (not published)")
            saw_404 = True
        except Exception as exc:  # noqa: BLE001 — try the next transport
            errors.append(f"{t}: {type(exc).__name__}: {exc}"[:160])
    if saw_404:
        raise FileNotFoundError(
            "a transport saw 404 (path may be wrong): "
            + " | ".join(errors))
    raise RuntimeError("all transports failed: " + " | ".join(errors))


_CACHE_MAX_AGE_S = 6 * 3600   # DOL appends in-quarter — hours-bounded


def _cache_dir() -> Path:
    d = os.environ.get("H1B_CACHE_DIR", "/tmp/h1b_quarter_cache")
    p = Path(d)
    p.mkdir(parents=True, exist_ok=True)
    return p


def _download_cached(url: str, transport: str, cfg,
                     timeout: int = 600) -> tuple[bytes, str]:
    """Run-#7 post-mortem: the workflow invokes this script once PER
    EMPLOYER, and each invocation re-downloaded every quarterly xlsx
    (83-251MB each; 9 employers x 6 quarters = ~54 downloads ~= 54 min
    — past the 45-min job timeout at employer #8, cancelling the run).
    A disk cache keyed by the quarter's file name shares ONE download
    across all invocations on the same runner (/tmp persists across
    the step's python processes; both the /sites/ and /media/ paths
    share the basename, so the alt-path retry cache-hits too).
    Bounded by mtime (6h) and size-verified; atomic rename keeps
    partial downloads out of the cache."""
    name = url.rsplit("/", 1)[-1]
    body_p = _cache_dir() / name
    meta_p = body_p.with_suffix(".xlsx.meta")
    if body_p.exists() and meta_p.exists():
        try:
            size = json.loads(meta_p.read_text(encoding="utf-8"))["size"]
            st = body_p.stat()
            if (time.time() - st.st_mtime < _CACHE_MAX_AGE_S
                    and st.st_size == size):
                return body_p.read_bytes(), "cache"
        except Exception:
            pass  # corrupt meta → fall through to a fresh download
    body, via = _download(url, transport, cfg, timeout)
    tmp = body_p.with_suffix(".xlsx.part")
    tmp.write_bytes(body)
    tmp.rename(body_p)  # atomic on the same filesystem
    meta_p.write_text(json.dumps({"size": len(body), "url": url}),
                      encoding="utf-8")
    return body, via


def _fetch_via(t: str, url: str, cfg, timeout: int) -> tuple[bytes, int]:
    import requests
    if t == "direct":
        r = requests.get(url, timeout=timeout, headers={
            "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                          "AppleWebKit/537.36 Chrome/128.0.0.0 "
                          "Safari/537.36"})
        if r.status_code == 404:
            raise FileNotFoundError(f"404 (quarter not published)")
        r.raise_for_status()
        return r.content, int(r.headers.get("Content-Length") or 0)
    if t == "impersonate":
        from curl_cffi import requests as creq
        r = creq.get(url, impersonate="chrome131", timeout=timeout)
        if r.status_code == 404:
            raise FileNotFoundError("404 (quarter not published)")
        r.raise_for_status()
        return r.content, int(r.headers.get("Content-Length") or 0)
    if t == "supabase":
        q = urlencode({"url": url, "mode": "raw"})
        r = requests.get(
            f"{cfg.supabase_proxy_url}?{q}",
            headers={"Authorization": f"Bearer "
                     f"{cfg.supabase_proxy_token}",
                     "x-region": "us-east-1"},
            timeout=timeout)
        if r.status_code == 404:
            raise FileNotFoundError("404 (quarter not published)")
        if r.status_code >= 400:
            raise RuntimeError(f"proxy HTTP {r.status_code}")
        return r.content, int(r.headers.get("Content-Length") or 0)
    raise ValueError(f"unknown transport {t!r}")


def _extract_rows(body: bytes, employer: str, source_file: str
                  ) -> tuple[list[dict], int]:
    """Filter the xlsx to employer-matching rows (streaming read-only
    mode — the files are too large for a full in-memory parse).

    `employer` is a comma-separated substring list (S12 multi-company:
    one company files under several legal names — "TENCENT AMERICA",
    "TENCENT AMERICA, INC."); a row matches when ANY substring hits
    (case-insensitive)."""
    import io
    import zipfile
    try:
        zf = zipfile.ZipFile(io.BytesIO(body))
    except zipfile.BadZipFile as exc:
        raise RuntimeError(f"not a valid xlsx (truncated body?): {exc}")
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(body), read_only=True,
                                data_only=True)
    ws = wb.active
    rows_out: list[dict] = []
    total = 0
    header: list[str] = []
    needles = tuple(n.strip().lower() for n in employer.split(",")
                    if n.strip())
    for row in ws.iter_rows(values_only=True):
        if not row:
            continue
        if not header:
            header = [str(c or "").strip() for c in row]
            continue
        total += 1
        rec = dict(zip(header, row))
        emp = str(rec.get("EMPLOYER_NAME") or "").lower()
        if not any(n in emp for n in needles):
            continue
        out = {"sourceFile": source_file}
        for col, field in _FIELDS:
            v = rec.get(col)
            out[field] = ("" if v is None
                          else v.isoformat() if hasattr(v, "isoformat")
                          else str(v))
        rows_out.append(out)
    wb.close()
    return rows_out, total


def _extract_rows_multi(body: bytes, pairs: list[tuple[str, str]],
                       source_file: str
                       ) -> dict[str, tuple[list[dict], int]]:
    """ONE streaming parse of the quarter's xlsx, every employer filter
    applied in the same pass (the --multi contract: parse cost is paid
    once, not once per employer — 9 employers x 6 quarters x 4.1M rows
    was ~100 min of re-parsing). Returns {label: (rows, total)}."""
    import io
    import zipfile
    try:
        zipfile.ZipFile(io.BytesIO(body))
    except zipfile.BadZipFile as exc:
        raise RuntimeError(f"not a valid xlsx (truncated body?): {exc}")
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(body), read_only=True,
                                data_only=True)
    ws = wb.active
    specs: dict[str, tuple[tuple[str, ...], list[dict]]] = {}
    for label, employer in pairs:
        needles = tuple(n.strip().lower() for n in employer.split(",")
                        if n.strip())
        specs[label] = (needles, [])
    total = 0
    header: list[str] = []
    for row in ws.iter_rows(values_only=True):
        if not row:
            continue
        if not header:
            header = [str(c or "").strip() for c in row]
            continue
        total += 1
        rec = dict(zip(header, row))
        emp = str(rec.get("EMPLOYER_NAME") or "").lower()
        if not emp:
            continue
        for label, (needles, out_rows) in specs.items():
            if any(n in emp for n in needles):
                out = {"sourceFile": source_file}
                for col, field in _FIELDS:
                    v = rec.get(col)
                    out[field] = ("" if v is None
                                  else v.isoformat()
                                  if hasattr(v, "isoformat")
                                  else str(v))
                out_rows.append(out)
    wb.close()
    return {label: (rows, total)
            for label, (needles, rows) in specs.items()}


def _existing_case_numbers(path: Path) -> set[str]:
    have: set[str] = set()
    if path.exists():
        for line in path.read_text(encoding="utf-8").split("\n"):
            line = line.strip()
            if not line:
                continue
            try:
                rec = json.loads(line)
            except json.JSONDecodeError:
                continue
            if rec.get("caseNumber"):
                have.add(rec["caseNumber"])
    return have


def _quarter_body(q: str, args, cfg
                  ) -> tuple[bytes | None, str, bool]:
    """Download one quarter (cache-aware, /media/ alt-path retry).
    Returns (body, via, failed); body=None means skip this quarter."""
    print(f"1b] {q}: downloading ({args.transport}) …", flush=True)
    try:
        body, via = _download_cached(_file_url(q), args.transport, cfg)
    except FileNotFoundError:
        # some quarters publish ONLY at /media/ (FY2026_Q3
        # live-measured: /sites/ 404s, /media/ serves) — one retry on
        # the alternate path before soft-skipping
        try:
            body, via = _download_cached(_file_url(q, alt=True),
                                         args.transport, cfg)
        except FileNotFoundError:
            print(f"  {q}: NOT PUBLISHED at either path — skipping "
                  f"(soft)", flush=True)
            return None, "", False
        except Exception as exc:  # noqa: BLE001
            print(f"  {q}: DOWNLOAD FAILED (alt path) — {exc}",
                  flush=True)
            return None, "", True
        print(f"  {q}: served from the /media/ path", flush=True)
    except Exception as exc:  # noqa: BLE001
        print(f"  {q}: DOWNLOAD FAILED — {exc}", flush=True)
        return None, "", True
    return body, via, False


def _parse_multi_spec(spec: str) -> list[tuple[str, str]]:
    pairs: list[tuple[str, str]] = []
    for chunk in spec.split(";"):
        chunk = chunk.strip()
        if not chunk:
            continue
        label, _, employer = chunk.partition(":")
        label, employer = label.strip(), employer.strip()
        if not label or not employer:
            raise ValueError(f"bad --multi pair {chunk!r} (expected "
                             "'label:employer-list')")
        pairs.append((label, employer))
    if not pairs:
        raise ValueError("--multi needs 'label:employer;label:employer'")
    return pairs


def _run_multi(args, cfg, quarters: list[str]) -> int:
    """The single-pass mode: per quarter ONE download (cached) + ONE
    parse + N employer filters; per label an independent append-only,
    case-number-deduped JSONL (the same contract as single mode)."""
    pairs = _parse_multi_spec(args.multi)
    out_paths: dict[str, Path] = {}
    haves: dict[str, set[str]] = {}
    files: dict[str, "object"] = {}
    added: dict[str, int] = {label: 0 for label, _ in pairs}
    employers = dict(pairs)
    try:
        for label, employer in pairs:
            p = (Path(args.out_dir) / label
                 ).with_suffix(".h1b_lca.jsonl")
            p.parent.mkdir(parents=True, exist_ok=True)
            out_paths[label] = p
            haves[label] = _existing_case_numbers(p)
            files[label] = open(p, "a", encoding="utf-8")
            print(f"1b] target: {p.name} "
                  f"({len(haves[label])} existing case numbers)",
                  flush=True)
        failed: list[str] = []
        for i, q in enumerate(quarters, 1):
            body, via, q_failed = _quarter_body(q, args, cfg)
            if q_failed:
                failed.append(q)
            if body is None:
                continue
            per_label = _extract_rows_multi(body, pairs, q)
            for label, (rows, total) in per_label.items():
                f = files[label]
                new = 0
                for rec in rows:
                    cn = rec.get("caseNumber") or ""
                    if not cn or cn in haves[label]:
                        continue
                    haves[label].add(cn)
                    f.write(json.dumps(rec, ensure_ascii=False) + "\n")
                    new += 1
                f.flush()
                added[label] += new
                print(f"  {q}: {len(body):,}B via {via}; {total:,} LCA "
                      f"rows, {len(rows)} {employers[label]} "
                      f"[{label}]; +{new} new (dedup case numbers)",
                      flush=True)
            if i < len(quarters):
                time.sleep(args.sleep)
    finally:
        for f in files.values():
            f.close()
    for label in added:
        print(f"1b] extract complete: +{added[label]} rows → "
              f"{out_paths[label]}", flush=True)
    if failed:
        print(f"1b] FAILED quarters (retry next run): "
              f"{', '.join(failed)}", file=sys.stderr)
        return 1
    return 0


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--quarters", default="",
                    help="comma-separated FYxxxx_Qn list (overrides "
                         "--recent)")
    ap.add_argument("--recent", type=int, default=0,
                    help="extract the N most recently published "
                         "quarters")
    ap.add_argument("--probe", action="store_true",
                    help="list published quarters + egress check, no "
                         "extraction")
    ap.add_argument("--employer", default="NVIDIA",
                    help="employer substring filter, comma-separated "
                         "list = OR (case-insensitive)")
    ap.add_argument("--label", default="nvidia_us_fulltime")
    ap.add_argument("--out-dir", default=str(
        REPO / "ingest" / "data" / "workday"))
    ap.add_argument("--transport", default="auto",
                    choices=["auto", "direct", "impersonate",
                             "supabase"])
    ap.add_argument("--sleep", type=float, default=2.0,
                    help="pause between file downloads (s)")
    ap.add_argument("--multi", default="",
                    help="single-pass multi-employer mode: "
                         "'label:employer-list' pairs joined by ';' "
                         "(employer lists keep their commas). Each "
                         "quarterly file is downloaded AND PARSED once, "
                         "every employer filter applied in the same "
                         "stream — the runs #7/#8 post-mortem: 9 "
                         "per-employer invocations re-parsed 4.1M rows "
                         "x9 (~100 min) and blew every timeout. "
                         "Overrides --label/--employer.")
    args = ap.parse_args()

    cfg = _load_env()
    quarters: list[str] = []
    if args.quarters:
        quarters = [q.strip().upper() for q in
                    args.quarters.split(",") if q.strip()]
    else:
        n = args.recent or (6 if not args.probe else 0)
        if n:
            quarters = _list_quarters(cfg)[-n:]
    if args.probe:
        published = _list_quarters(cfg)
        print(f"published quarters ({len(published)}): "
              f"{', '.join(published[-8:])}"
              + (" …" if len(published) > 8 else ""))
        for q in (quarters or published[-1:]):
            url = _file_url(q)
            try:
                import requests
                qq = urlencode({"url": url, "mode": "raw"})
                r = requests.get(
                    f"{cfg.supabase_proxy_url}?{qq}",
                    headers={"Authorization": f"Bearer "
                             f"{cfg.supabase_proxy_token}",
                             "x-region": "us-east-1"},
                    timeout=60, stream=True)
                print(f"  {q}: HTTP {r.status_code}, "
                      f"content-length {r.headers.get('content-length')}")
                r.close()
            except Exception as exc:  # noqa: BLE001
                print(f"  {q}: probe failed: {exc}")
        return 0
    if not quarters:
        print("nothing to do: pass --quarters or --recent N")
        return 2
    if args.multi:
        return _run_multi(args, cfg, quarters)

    out_path = (Path(args.out_dir) / args.label
                ).with_suffix(".h1b_lca.jsonl")
    out_path.parent.mkdir(parents=True, exist_ok=True)
    have = _existing_case_numbers(out_path)
    print(f"[h1b] target: {out_path.name} "
          f"({len(have)} existing case numbers)")
    failed: list[str] = []
    added = 0
    with open(out_path, "a", encoding="utf-8") as f:
        for i, q in enumerate(quarters, 1):
            print(f"[h1b] [{i}/{len(quarters)}] {q}: downloading "
                  f"({args.transport}) …", flush=True)
            body = None
            via = ""
            try:
                body, via = _download_cached(_file_url(q),
                                           args.transport, cfg)
            except FileNotFoundError:
                # some quarters publish ONLY at /media/ (FY2026_Q3
                # live-measured: /sites/ 404s, /media/ serves) — one
                # retry on the alternate path before soft-skipping
                try:
                    body, via = _download_cached(_file_url(q, alt=True),
                                               args.transport, cfg)
                except FileNotFoundError:
                    print(f"  {q}: NOT PUBLISHED at either path — "
                          f"skipping (soft)", flush=True)
                    continue
                except Exception as exc:  # noqa: BLE001
                    print(f"  {q}: DOWNLOAD FAILED (alt path) — {exc}",
                          flush=True)
                    failed.append(q)
                    continue
                print(f"  {q}: served from the /media/ path", flush=True)
            except Exception as exc:  # noqa: BLE001
                print(f"  {q}: DOWNLOAD FAILED — {exc}", flush=True)
                failed.append(q)
                continue
            if body is None:
                continue
            rows, total = _extract_rows(body, args.employer, q)
            new = 0
            for rec in rows:
                cn = rec.get("caseNumber") or ""
                if not cn or cn in have:
                    continue
                have.add(cn)
                f.write(json.dumps(rec, ensure_ascii=False) + "\n")
                new += 1
            f.flush()
            added += new
            print(f"  {q}: {len(body):,}B via {via}; {total:,} LCA "
                  f"rows, {len(rows)} {args.employer}; +{new} new "
                  f"(dedup case numbers)", flush=True)
            if i < len(quarters):
                time.sleep(args.sleep)
    print(f"[h1b] extract complete: +{added} rows → {out_path}",
          flush=True)
    if failed:
        print(f"[h1b] FAILED quarters (retry next run): "
              f"{', '.join(failed)}", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
